import pandas as pd
import sqlite3
import os

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE_FILE = os.path.join(SCRIPT_DIR, "Database.xlsx")
SQLITE_DB = os.path.join(SCRIPT_DIR, "courses.db")

def setup_database():
    if os.path.exists(SQLITE_DB):
        return

    # Read Excel with correct header row (row 3) and extract hyperlinks from 'Courses' column
    import openpyxl
    wb = openpyxl.load_workbook(DATABASE_FILE, data_only=True)
    ws = wb.active
    # Always use column B (index 2, 1-based) for course links
    header_row = 4  # 1-based (row 4 in Excel, header=3 in pandas)
    courses_col_idx = 2  # Column B
    # Read with pandas as before
    df = pd.read_excel(DATABASE_FILE, header=3)
    # Extract hyperlinks from column B for all data rows
    course_links = []
    for i, row in enumerate(ws.iter_rows(min_row=header_row+1, max_row=ws.max_row), 1):
        cell = row[1]  # Column B, 0-based index
        if cell and cell.hyperlink:
            course_links.append(cell.hyperlink.target)
        else:
            course_links.append("")
    df['course_link'] = course_links[:len(df)]

    column_mapping = {
        "COURSE NAME": "course_name",
        "COURSE CODE": "course_code",
        "MAJOR": "major",
        "CREDITS": "credits",
        "PASS PERCENTAGE ": "pass_rate",
        "STUDENT DIFFICULTY RATING": "difficulty",
        "COURSE DESCRIPTION": "description",
        "INTENDED LEARNINGS": "learning_outcomes",
        "STUDENT COURSE RATING": "rating",
        "MANDATORY/ELECTIVE": "mandatory_elective",
        "OFFERED IN PERIOD": "offered_period",
        "EXAMINATION": "examination",
        "GRADING CRITERIA": "grading_criteria",
        "PREREQUISITE": "prerequisite",
        "courses": "course_link"
    }
    
    df = df.rename(columns=column_mapping)
    
    # Convert pass_rate from percentage to decimal (e.g., 80 = 0.8)
    if "pass_rate" in df.columns:
        df["pass_rate"] = pd.to_numeric(df["pass_rate"], errors='coerce') / 100
    
    # Ensure numeric columns
    df["credits"] = pd.to_numeric(df["credits"], errors='coerce')
    df["difficulty"] = pd.to_numeric(df["difficulty"], errors='coerce')
    
    # Add a 'topics' field (basic keyword tagging for now)
    def tag_topics(row):
        # ai_keywords removed; topic_keywords covers all topics
        topic_keywords = {
                "AI": ["ai", "artificial intelligence", "machine learning", "deep learning", "neural network", "generative ai", "nlp", "computer vision"],
                "Programming": ["programming", "software engineering", "coding", "python", "java", "c++", "development", "software"],
                "Data Science": ["data science", "data analysis", "statistics", "data mining", "big data", "analytics"],
                "Sustainability": ["sustainability", "ecology", "environment", "green", "renewable"],
                "Ethics": ["ethics", "ethical", "responsibility", "fairness", "bias"],
                "Design": ["design", "ux", "user experience", "interface", "visualization", "graphics"],
                "HCI": ["human-computer interaction", "hci", "usability", "interaction"],
                "XR": ["extended reality", "virtual reality", "augmented reality", "xr", "vr", "ar"],
                "Game Development": ["game", "gaming", "game design", "game development"],
                "Project": ["project", "project work", "project-based"],
                "Seminar": ["seminar", "seminar-based"],
                "Lab": ["lab", "laboratory", "hands-on"],
                "Visualization": ["visualization", "visualisation", "graphics", "image processing", "video processing"]
            }
        import re
        text = f"{row.get('course_name','')} {row.get('description','')} {row.get('learning_outcomes','')} {row.get('prerequisite','')}".lower()
        tags = set()
        for topic, keywords in topic_keywords.items():
            for kw in keywords:
                if ' ' in kw:
                    if kw in text:
                        tags.add(topic)
                        break
                else:
                    if re.search(r'\\b' + re.escape(kw) + r'\\b', text):
                        tags.add(topic)
                        break
        return ",".join(sorted(tags))

    df["topics"] = df.apply(tag_topics, axis=1)

    # Add a short_description field (first 30 words of description)
    def short_desc(text):
        if not isinstance(text, str):
            return ""
        words = text.split()
        return " ".join(words[:30]) + ("..." if len(words) > 30 else "")
    df["short_description"] = df["description"].apply(short_desc)

    columns_to_keep = ["course_code", "course_name", "major", "credits", "pass_rate", 
                       "difficulty", "description", "short_description", "learning_outcomes", "rating",
                       "mandatory_elective", "offered_period", "examination", "grading_criteria", "topics", "prerequisite", "course_link"]
    df = df[[col for col in columns_to_keep if col in df.columns]]

    conn = sqlite3.connect(SQLITE_DB)
    df.to_sql("courses", conn, if_exists="replace", index=False)
    conn.close()
